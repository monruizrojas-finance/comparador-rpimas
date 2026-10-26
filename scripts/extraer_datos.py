#!/usr/bin/env python3
"""
Extrae del Informe Trimestral (Seguros Generales):
  - Prima Directa por compañía y ramo  (hoja "Gen. Prima Directa")
  - Estado de Resultados por compañía  (hoja "Gen. EERR")
y genera data/datos.js para el sitio.

Uso (uno o más informes; cada uno queda como un período en el sitio):
    pip install openpyxl
    python scripts/extraer_datos.py "Informe Trimestral 2026.06 - Web.xlsx" "Informe Trimestral 2025.12 - Web.xlsx"
"""
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")

HOJA_PRIMA = "Gen. Prima Directa"
HOJA_EERR = "Gen. EERR"
FILA_ENCABEZADO = 11
COL_INICIO = 3  # columna C
# Cuentas que restan dentro de su subtotal aunque la columna "Signos" del informe diga +1
RESTAN = {"5.31.33.30", "5.31.33.40"}  # Depreciación inversiones, Gastos de gestión


def limpiar(texto):
    return re.sub(r"\s+", " ", str(texto)).strip()


def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else 0.0


def leer_companias(ws):
    """Compañías desde la columna C hasta la primera vacía + columnas Mercado / Signos."""
    companias, col = [], COL_INICIO
    while ws.cell(FILA_ENCABEZADO, col).value not in (None, ""):
        companias.append((col, limpiar(ws.cell(FILA_ENCABEZADO, col).value)))
        col += 1
    extra = {}
    for c in range(col, ws.max_column + 1):
        v = ws.cell(FILA_ENCABEZADO, c).value
        if v is not None:
            extra.setdefault(limpiar(v), c)
    return companias, extra


def leer_prima(ws):
    companias, extra = leer_companias(ws)
    col_mercado = extra["Mercado"]
    ramos, fila_total = [], None
    for r in range(FILA_ENCABEZADO + 1, ws.max_row + 1):
        etiqueta = ws.cell(r, 2).value
        if etiqueta is None:
            continue
        etiqueta = limpiar(etiqueta)
        if etiqueta == "Total":
            fila_total = r
            break
        m = re.match(r"^([\d.]+)\s+(.*)$", etiqueta)
        if not m:
            continue
        ramos.append({
            "codigo": m.group(1).rstrip("."),
            "nombre": m.group(2),
            "mercado": num(ws.cell(r, col_mercado).value),
            "valores": [num(ws.cell(r, c).value) for c, _ in companias],
        })
    return {
        "companias": [n for _, n in companias],
        "totales": [num(ws.cell(fila_total, c).value) for c, _ in companias],
        "mercado": num(ws.cell(fila_total, col_mercado).value),
        "mercado_anterior": num(ws.cell(fila_total, extra["Mercado " + fecha_anterior(ws)]).value),
        "fecha_anterior": fecha_anterior(ws),
        "ramos": ramos,
    }


def fecha_anterior(ws):
    """Fecha del mismo trimestre del año anterior, p. ej. 30.06.2026 -> 30.06.2025"""
    d, m, a = limpiar(ws["B9"].value).split(" al ")[-1].split(".")
    return f"{d}.{m}.{int(a) - 1}"


def nivel_cuenta(codigo):
    """5.31.00.00 -> 0 (total) · 5.31.10.00 -> 1 · 5.31.11.00 -> 2 · 5.31.11.10 -> 3"""
    p = codigo.split(".")
    if p[2] == "00":
        return 0
    if p[3] != "00":
        return 3
    return 1 if p[2].endswith("0") else 2


def leer_eerr(ws, companias_esperadas):
    companias, extra = leer_companias(ws)
    nombres = [n for _, n in companias]
    if nombres != companias_esperadas:
        raise SystemExit("Las compañías del EERR no coinciden con las de Prima Directa")
    col_mercado, col_signo = extra["Mercado"], extra["Signos"]
    cuentas = []
    for r in range(FILA_ENCABEZADO + 1, ws.max_row + 1):
        etiqueta = ws.cell(r, 2).value
        if etiqueta is None:
            break
        m = re.match(r"^(\d\.\d{2}\.\d{2}\.\d{2})\s+(.*)$", limpiar(etiqueta))
        if not m:
            break
        signo = ws.cell(r, col_signo).value or 1  # el informe guarda montos positivos + signo contable
        if m.group(1) in RESTAN:
            signo = -1
        cuentas.append({
            "codigo": m.group(1),
            "nombre": m.group(2).capitalize() if m.group(2).istitle() else m.group(2),
            "nivel": nivel_cuenta(m.group(1)),
            "mercado": num(num(ws.cell(r, col_mercado).value) * signo),
            "valores": [num(num(ws.cell(r, c).value) * signo) for c, _ in companias],
        })
    return cuentas


def leer_periodo(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    subtitulo = limpiar(wb[HOJA_PRIMA]["B9"].value)  # "Cifras en UF al 30.06.2026"
    prima = leer_prima(wb[HOJA_PRIMA])
    eerr = leer_eerr(wb[HOJA_EERR], prima["companias"])
    periodo = {
        "fecha": subtitulo.split(" al ")[-1],
        "unidad": "UF",
        "fuente": Path(ruta).name,
        "companias": prima["companias"],
        "prima": {k: prima[k] for k in ("totales", "mercado", "mercado_anterior", "fecha_anterior", "ramos")},
        "eerr": eerr,
    }
    print(f"{periodo['fecha']}: {len(periodo['companias'])} compañías · "
          f"{len(prima['ramos'])} ramos · {len(eerr)} cuentas de EERR")
    return periodo


def orden_fecha(p):
    d, m, a = p["fecha"].split(".")
    return (int(a), int(m), int(d))


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    periodos = [leer_periodo(r) for r in sys.argv[1:]]
    periodos.sort(key=orden_fecha, reverse=True)  # el más reciente primero
    fechas = [p["fecha"] for p in periodos]
    if len(set(fechas)) != len(fechas):
        sys.exit("Hay dos informes con la misma fecha de corte")

    # Se guarda como .js para que el sitio funcione también abriendo index.html con doble clic
    destino = Path(__file__).resolve().parent.parent / "data" / "datos.js"
    contenido = json.dumps({"periodos": periodos}, ensure_ascii=False, separators=(",", ":"))
    destino.write_text(f"window.DATOS = {contenido};\n", encoding="utf-8")
    print(f"Escrito {destino} con {len(periodos)} períodos: {', '.join(fechas)}")


if __name__ == "__main__":
    main()
